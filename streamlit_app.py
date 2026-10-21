"""Rumaila Pattern Waterflood | upload-driven Streamlit dashboard.

Run locally: streamlit run streamlit_app.py
GitHub/Streamlit Cloud: deploy with streamlit_app.py + requirements.txt.
No field data is bundled in this source code.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Optional

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

st.set_page_config(page_title="Pattern Waterflood | Rumaila", page_icon="💧", layout="wide", initial_sidebar_state="expanded")

# -----------------------------------------------------------------------------
# Input definitions — these are matched without regard to case or punctuation.
# -----------------------------------------------------------------------------
FIELDS = {
    "Pattern": ["Pattern Name", "Pattern", "Pattern ID"],
    "Reservoir": ["Reservoir", "Reservoir Name", "Resevoir", "Resevoir Name", "Reservior", "Reservior Name", "Reservoir Type"],
    "Manifold": ["Manifold", "Manifold Name", "Injection Manifold", "Production Manifold", "Manifold ID", "Manifold Code"],
    "Oil": ["Produced Oil (stb)", "Produced Oil", "Oil Production (stb)", "Oil Volume"],
    "Water": ["Produced Water (stb)", "Produced Water", "Water Production (stb)"],
    "Injection": ["Injected Water (stb)", "Injected Water", "Water Injection (stb)", "Water Injected"],
    "VRR": ["VRR", "Voidage Replacement Ratio"],
    "Res Pressure": ["Average of Res P", "Average Reservoir Pressure", "Reservoir Pressure"],
    "Psat": ["Average of P Sat", "Average Saturation Pressure", "Bubble Point Pressure"],
    "Delta P": ["Delta P from P Sat", "Delta P from Psat", "Pressure Margin"],
    "FBHP Trend": ["Avg. FBHP Trend (psi/day)", "FBHP Trend (psi/day)", "Average FBHP Trend"],
    "Watercut Input": ["Watercut %", "Water Cut", "Watercut", "Water Cut %"],
    "Agreed Target": ["Agreed Target (stb)", "Agreed Target"],
    "RMT Target": ["RMT Target (stb)", "RMT Target"],
    "Actual Supported Oil": ["Oil Supported (stb)", "Supported Oil (stb)", "Actual Supported Oil (stb)", "Incremental Oil (stb)"],
    "Support Fraction": ["Oil Support Fraction", "Supported Fraction", "Support Factor", "Support Fraction"],
}
REQUIRED = ["Pattern", "Oil", "Water", "Injection"]
NUMERIC_FIELDS = [k for k in FIELDS if k not in ("Pattern", "Reservoir", "Manifold")]
VOL_STB_TO_M3 = 0.158987294928
UNIT_OPTIONS = {
    "stb": (1.0, "stb"),
    "Mstb": (1e-3, "Mstb"),
    "MMstb": (1e-6, "MMstb"),
    "m³": (VOL_STB_TO_M3, "m³"),
    "10³ m³": (VOL_STB_TO_M3 / 1_000, "10³ m³"),
}
DEFAULT_CPS_COLORS = {
    "CPS1": "#1677B9", "CPS2": "#DA6D24", "CPS3": "#159579",
    "CPS4": "#8E62A6", "CPS5": "#C43B5F", "CPS6": "#807026",
    "CPS7": "#18A9B4", "Unassigned": "#8A929C",
}


@dataclass
class Design:
    font: str = "Arial"
    text: str = "#182536"
    paper: str = "#FFFFFF"
    plot: str = "#FFFFFF"
    grid: str = "#E7EDF5"
    title_size: int = 19
    axis_size: int = 14
    tick_size: int = 11
    legend_size: int = 11
    line_width: float = 2.5
    marker_size: int = 13
    opacity: float = 0.83
    height: int = 570
    legend_pos: str = "Bottom"
    legend_x: float = 0.5
    legend_y: float = -0.27
    legend_anchor: str = "center"
    grid_on: bool = True


def keyify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def col_lookup(frame: pd.DataFrame, aliases: list[str]) -> Optional[str]:
    dic = {keyify(x): x for x in frame.columns}
    for term in aliases:
        if keyify(term) in dic:
            return dic[keyify(term)]
    return None


# Canonical labels used only for recognising a misnamed reservoir column.
# Other reservoir names are preserved when the workbook contains a valid header.
RESERVOIR_LABELS = {
    "mishrif": "Mishrif",
    "uppershale": "Upper Shale",
    "mainpay": "Main Pay",
}


def canonical_reservoir(value: object) -> str:
    if pd.isna(value) or not str(value).strip():
        return "Unspecified"
    original = str(value).strip()
    return RESERVOIR_LABELS.get(keyify(original), original)


def detect_reservoir_column(frame: pd.DataFrame) -> Optional[str]:
    """Find a mislabeled reservoir column by values, not just the header.

    Some Rumaila files have B1='Mishrif' instead of B1='Reservoir', although
    the cells below contain Mishrif, Upper Shale and Main Pay. Requiring a
    strong concentration of *known reservoir labels* avoids mistaking Pattern
    names or other text columns for a reservoir.
    """
    candidates: list[tuple[float, int, str]] = []
    for col in frame.columns:
        if keyify(col) in {"pattern", "patternname", "patternid"}:
            continue
        values = frame[col].dropna()
        if values.empty:
            continue
        strings = values.astype(str).str.strip()
        strings = strings.loc[strings.ne("")]
        if len(strings) < 3:
            continue
        recognized = strings.map(lambda v: keyify(v) in RESERVOIR_LABELS)
        confidence = float(recognized.mean())
        known_kinds = strings.loc[recognized].map(keyify).nunique()
        is_reservoir_header = keyify(col) in RESERVOIR_LABELS
        if confidence >= 0.80 and (is_reservoir_header or known_kinds >= 2):
            candidates.append((confidence, int(recognized.sum()), str(col)))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def to_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    cleaned = (series.astype("string").str.replace(",", "", regex=False)
               .str.replace("%", "", regex=False)
               .str.replace("−", "-", regex=False).str.strip())
    return pd.to_numeric(cleaned, errors="coerce")


def get_cps(pattern: str) -> tuple[str, str]:
    text = str(pattern).upper()
    match = re.search(r"(?<![A-Z0-9])CPS[\s_-]*(\d+)(?!\d)", text)
    if match:
        return f"CPS{int(match.group(1))}", ""
    # One known typo in the October source: B-CSP4-i9-R-856.
    typo = re.search(r"(?<![A-Z0-9])CSP[\s_-]*(\d+)(?!\d)", text)
    if typo:
        return f"CPS{int(typo.group(1))}", "CSP/CPS spelling normalised; check source"
    return "Unassigned", "CPS not found in pattern name"


@st.cache_data(show_spinner=False)
def sheet_names(file_bytes: bytes, extension: str) -> list[str]:
    if extension == "csv":
        return ["CSV"]
    excel = pd.ExcelFile(io.BytesIO(file_bytes))
    return excel.sheet_names


@st.cache_data(show_spinner=False)
def read_sheet(file_bytes: bytes, extension: str, name: str) -> pd.DataFrame:
    if extension == "csv":
        return pd.read_csv(io.BytesIO(file_bytes))
    return pd.read_excel(io.BytesIO(file_bytes), sheet_name=name)


def clean_manifold(value: object) -> str:
    """Normalise explicit Manifold identifiers without inferring them from Pattern names."""
    if pd.isna(value):
        return "Unmapped"
    raw = str(value).strip()
    if not raw or raw.lower() in ("nan", "none", "unmapped", "unknown", "n/a", "-", "tbc", "not assigned"):
        return "Unmapped"
    # Operational comments are not real manifolds; keep them visible in source QC.
    if any(phrase in raw.lower() for phrase in ("not converted", "not an injector", "to be converted", "no manifold")):
        return "Unmapped"
    raw = re.sub(r"\s*[/\\]\s*", "/", raw.strip()).upper()
    raw = re.sub(r"\s+", " ", raw)
    # e.g., CPS-2 / FM09, CPS2/FM09, cps-2/SM01 => CPS-2/FM09
    match = re.fullmatch(r"CPS[\s_-]*(\d+)\s*/\s*([A-Z0-9-]+)", raw)
    if match:
        return f"CPS-{int(match.group(1))}/{match.group(2)}"
    return raw


def manifold_cps(name: str) -> Optional[str]:
    """CPS declared by the Manifold ID; no Pattern-name guess when present."""
    m = re.match(r"^CPS[\s_-]*(\d+)\s*/", str(name), flags=re.I)
    return f"CPS{int(m.group(1))}" if m else None


def manifold_rollup(frame: pd.DataFrame, keys: Optional[list[str]] = None) -> pd.DataFrame:
    """Group by physical Manifold ID, which already includes CPS e.g. CPS-2/FM09.

    Pattern CPS may disagree with the supplied Manifold CPS. Do not split the
    *same* Manifold into multiple points on account of that mismatch.
    """
    keys = ["Manifold"] + ([k for k in (keys or []) if k != "Manifold"])
    result = rollup(frame, keys)
    if result.empty:
        return result
    pattern_cps_sets = frame.groupby("Manifold")["CPS"].agg(lambda ss: sorted(set(ss.dropna())))
    result["CPS"] = result["Manifold"].map(
        lambda mn: manifold_cps(mn) or (
            pattern_cps_sets[mn][0] if len(pattern_cps_sets[mn]) == 1 else "Mixed"
        )
    )
    return result


def normalise_vow(frame: pd.DataFrame, reference: float) -> pd.DataFrame:
    """Display-only 0–1 score. Never overwrite the observed volume ratio."""
    out = frame.copy()
    if out.empty:
        out["VoW Normalized"] = pd.Series(dtype=float)
        return out
    raw = out["VoW Technical"]
    out["VoW Normalized"] = (raw / reference).clip(lower=0, upper=1)
    return out


def vow_normalization_reference(frame: pd.DataFrame) -> float:
    """Current selected group reference, NEVER pooled over different levels.

    Example: when CPS is selected, max CPS ratio becomes 1 even if a
    Pattern elsewhere has a much higher raw VoW. Calculated on the
    *currently filtered* rows at the level actually displayed.
    """
    if frame.empty or "VoW Technical" not in frame.columns:
        return 1.0
    vals = pd.to_numeric(frame["VoW Technical"], errors="coerce")
    vals = vals.loc[np.isfinite(vals) & vals.gt(0)]
    return float(vals.max()) if len(vals) else 1.0


def map_manifolds(frame: pd.DataFrame, mapping: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Safely apply one Manifold per Pattern+Reservoir; never guess from a name.

    Supports a Pattern-only map only when that Pattern is unambiguous across
    reservoirs. Conflicting assignments stay unmapped with a QC warning.
    """
    df = frame.copy()
    notes: list[str] = []
    pc = col_lookup(mapping, ["Pattern", "Pattern Name", "Pattern ID"])
    mc = col_lookup(mapping, FIELDS["Manifold"])
    rc = col_lookup(mapping, FIELDS["Reservoir"])
    if not pc or not mc:
        return df, ["Mapping requires Pattern (or Pattern Name) and Manifold columns."]
    cols = [pc, mc] + ([rc] if rc else [])
    ref = mapping[cols].copy()
    ref.columns = ["Pattern", "Manifold"] + (["Reservoir"] if rc else [])
    ref["Pattern"] = ref["Pattern"].fillna("").astype(str).str.strip()
    ref["Manifold"] = ref["Manifold"].map(clean_manifold)
    ref = ref.loc[ref["Pattern"].ne("") & ref["Manifold"].ne("Unmapped")].copy()
    if ref.empty:
        return df, ["No assigned Manifold entries found in mapping; cross-plot requires a mapping."]
    df["_map_key"] = df["Pattern"].str.casefold()
    ref["_map_key"] = ref["Pattern"].str.casefold()
    join = ["_map_key"]
    if rc:
        df["_res_key"] = df["Reservoir"].str.casefold()
        ref["_res_key"] = ref["Reservoir"].fillna("").astype(str).str.strip().map(canonical_reservoir).str.casefold()
        join.append("_res_key")
    else:
        amb = df.groupby("_map_key")["Reservoir"].nunique()
        ambiguous = set(amb.loc[amb.gt(1)].index)
        if ambiguous:
            ref = ref.loc[~ref["_map_key"].isin(ambiguous)].copy()
            notes.append(f"{len(ambiguous)} Pattern names span multiple reservoirs; include a Reservoir column in the mapping to disambiguate them.")
    conflicts = ref.groupby(join, dropna=False)["Manifold"].nunique()
    conflict_keys = conflicts.loc[conflicts.gt(1)]
    if len(conflict_keys):
        notes.append(f"{len(conflict_keys)} conflicting Pattern→Manifold assignments ignored. Each Pattern needs one assignment; split multi-Manifold allocation upstream.")
        bad = ref.set_index(join).index.isin(conflict_keys.index)
        ref = ref.loc[~bad]
    ref = ref.drop_duplicates(join, keep="first")
    if ref.empty:
        return frame.copy(), notes
    assigned = df.merge(ref[join + ["Manifold"]].rename(columns={"Manifold": "_assigned"}), on=join, how="left", sort=False)
    changed = assigned["_assigned"].notna()
    assigned.loc[changed, "Manifold"] = assigned.loc[changed, "_assigned"]
    assigned = assigned.drop(columns=["_map_key", "_res_key", "_assigned"], errors="ignore")
    notes.append(f"{int(changed.sum())} source rows matched to a Manifold.")
    return assigned, notes


def display_volume(values: pd.Series, factor: float, months: float, mode: str, days: int) -> pd.Series:
    divisor = days if mode == "Average per day" else (months if mode == "Average per month" else 1)
    return values * factor / divisor


def attach_display(frame: pd.DataFrame, factor: float, months: float, mode: str, days: int) -> pd.DataFrame:
    frame = frame.copy()
    for original, name in (("Oil", "Oil Display"), ("Water", "Water Display"),
                           ("Injection", "Injection Display"), ("Supported", "Supported Display")):
        frame[name] = display_volume(frame[original], factor, months, mode, days)
    frame["Supported mbd"] = frame["Supported"] / (days * 1000.0)
    frame["Injection mbd"] = frame["Injection"] / (days * 1000.0)
    return frame


def standardise(raw: pd.DataFrame) -> tuple[pd.DataFrame, list[str], dict[str, str]]:
    # Detect before discarding unnamed columns: the reservoir header is
    # sometimes overwritten (e.g., B1='Mishrif') or left blank in Excel.
    matched = {k: col_lookup(raw, aliases) for k, aliases in FIELDS.items()}
    if matched["Reservoir"] is None:
        matched["Reservoir"] = detect_reservoir_column(raw)
    missing = [k for k in REQUIRED if matched[k] is None]
    if missing:
        return pd.DataFrame(), missing, {k: v for k, v in matched.items() if v}
    df = pd.DataFrame(index=raw.index)
    for key in FIELDS:
        if matched[key]:
            df[key] = raw[matched[key]]
        else:
            # Missing Reservoir must never silently become Mishrif.
            df[key] = "Unspecified" if key == "Reservoir" else ("Unmapped" if key == "Manifold" else np.nan)
    df["Pattern"] = df["Pattern"].astype("string").str.strip()
    df = df[df["Pattern"].notna() & df["Pattern"].ne("")].copy()
    df["Reservoir"] = df["Reservoir"].map(canonical_reservoir)
    df["Manifold"] = df["Manifold"].map(clean_manifold)
    for key in NUMERIC_FIELDS:
        df[key] = to_numeric(df[key])
    df["Manifold CPS"] = df["Manifold"].map(manifold_cps)
    # User files may store fractions [0,1] or percentage points [0,100].
    for pct_col in ["Watercut Input", "Support Fraction"]:
        non_null = df[pct_col].dropna()
        if len(non_null) and float(non_null.quantile(0.95)) > 1.0:
            df[pct_col] = df[pct_col] / 100.0
    cps = df["Pattern"].map(get_cps)
    df["CPS"] = cps.map(lambda pair: pair[0])
    df["CPS QC"] = cps.map(lambda pair: pair[1])
    df["Manifold CPS mismatch"] = (
        df["Manifold CPS"].notna() & df["CPS"].ne("Unassigned") & df["Manifold CPS"].ne(df["CPS"])
    )
    df["Source Row"] = df.index.astype(int) + 2
    df["Liquid"] = df["Oil"].fillna(0) + df["Water"].fillna(0)
    df["Watercut Calculated"] = np.where(df["Liquid"] > 0, df["Water"].fillna(0) / df["Liquid"], np.nan)
    df["Oil vs Agreed"] = df["Oil"] - df["Agreed Target"]
    df["Oil vs RMT"] = df["Oil"] - df["RMT Target"]
    return df.reset_index(drop=True), [], {k: v for k, v in matched.items() if v}


def support_scenario(frame: pd.DataFrame, method: str, share: float) -> pd.DataFrame:
    df = frame.copy()
    if method == "Excel: Oil Supported":
        df["Supported"] = df["Actual Supported Oil"]
    elif method == "Excel: Support Fraction":
        df["Supported"] = df["Oil"] * df["Support Fraction"]
    else:
        df["Supported"] = df["Oil"] * share / 100.0
    return df


def safe_ratio(numerator: pd.Series, denominator: pd.Series) -> pd.Series:
    return numerator.div(denominator.where(denominator > 0)).replace([np.inf, -np.inf], np.nan)


def rollup(frame: pd.DataFrame, keys: list[str]) -> pd.DataFrame:
    """CRITICAL: aggregate numerator/denominator FIRST, then calculate VoW."""
    if frame.empty:
        return pd.DataFrame()
    sum_cols = ["Oil", "Water", "Injection", "Supported", "Agreed Target", "RMT Target"]
    aggs: dict[str, object] = {c: (lambda x: x.sum(min_count=1)) for c in sum_cols}
    for col in ("VRR", "Res Pressure", "Psat", "Delta P", "FBHP Trend"):
        aggs[col] = "mean"  # Informational only; NOT a recomputed reservoir-condition VRR.
    if "Pattern" not in keys:
        aggs["Pattern"] = "nunique"
    result = frame.groupby(keys, dropna=False, as_index=False).agg(aggs)
    if "Pattern" not in keys:
        result = result.rename(columns={"Pattern": "Patterns"})
    else:
        result["Patterns"] = 1
    result["Liquid"] = result["Oil"].fillna(0) + result["Water"].fillna(0)
    result["Watercut"] = safe_ratio(result["Water"], result["Liquid"])
    result["VoW Technical"] = safe_ratio(result["Supported"], result["Injection"])
    # Raw ratio is kept for engineering QC; 0–1 score added later for charts.
    result["Production vs Agreed"] = result["Oil"] - result["Agreed Target"]
    result["Production vs RMT"] = result["Oil"] - result["RMT Target"]
    result["VoW Valid"] = result["Injection"].gt(0) & result["Supported"].notna()
    return result


def cps_sort(cps: str) -> tuple[int, str]:
    m = re.search(r"\d+", str(cps))
    return (int(m.group()) if m else 1000, str(cps))


def pretty(value: float, digits: int = 2) -> str:
    return "—" if pd.isna(value) else f"{value:,.{digits}f}"


def apply_style(fig: go.Figure, design: Design, title: str, x_label: str = "", y_label: str = "", *, show_legend: bool = True, height: Optional[int] = None) -> go.Figure:
    if design.legend_pos == "Hidden" or not show_legend:
        legend = dict()
        top_margin, bottom_margin, right_margin = 95, 100, 75
    elif design.legend_pos == "Top":
        legend = dict(orientation="h", x=0.5, y=1.14, xanchor="center", yanchor="bottom")
        top_margin, bottom_margin, right_margin = 135, 95, 70
    elif design.legend_pos == "Right":
        legend = dict(orientation="v", x=1.02, y=1, xanchor="left", yanchor="top")
        top_margin, bottom_margin, right_margin = 95, 100, 175
    elif design.legend_pos == "Custom":
        legend = dict(orientation="h", x=design.legend_x, y=design.legend_y, xanchor=design.legend_anchor, yanchor="middle")
        top_margin, bottom_margin, right_margin = 130, 150, 95
    else:
        legend = dict(orientation="h", x=0.5, y=-0.28, xanchor="center", yanchor="top")
        top_margin, bottom_margin, right_margin = 95, 165, 75
    legend.update(font=dict(size=design.legend_size, color=design.text), bgcolor="rgba(0,0,0,0)")
    fig.update_layout(
        title=dict(text=f"<b>{title}</b>", x=0.01, y=0.98, xanchor="left", yanchor="top", font=dict(size=design.title_size, color=design.text)),
        font=dict(family=design.font, color=design.text, size=design.tick_size),
        plot_bgcolor=design.plot, paper_bgcolor=design.paper,
        height=height or design.height, hovermode="closest",
        margin=dict(l=90, r=right_margin, b=bottom_margin, t=top_margin),
        legend=legend, showlegend=show_legend and design.legend_pos != "Hidden",
        hoverlabel=dict(font_size=max(11, design.tick_size), font_family=design.font),
    )
    for axis, label in (("xaxes", x_label), ("yaxes", y_label)):
        kwargs = {
            "title_text": label,
            "title_font": dict(size=design.axis_size, color=design.text),
            "title_standoff": 18,
            "tickfont": dict(size=design.tick_size, color=design.text),
            "showgrid": design.grid_on,
            "gridcolor": design.grid,
            "gridwidth": 1,
            "zeroline": False,
            "linecolor": design.grid,
            "showline": True,
            "automargin": True,
        }
        getattr(fig, f"update_{axis}")(**kwargs)
    fig.update_traces(selector=dict(type="scatter"), line=dict(width=design.line_width))
    return fig


def plot(fig: go.Figure, key: str):
    st.plotly_chart(fig, use_container_width=True, key=key,
                    config={"displaylogo": False, "scrollZoom": False, "toImageButtonOptions": {"format": "png", "filename": key, "scale": 2}})


def scale_volume(x: pd.Series, factor: float, months: float, monthly: bool) -> pd.Series:
    return x * factor / (months if monthly else 1)


def file_download_excel(patterns: pd.DataFrame, cps: pd.DataFrame, manifold: pd.DataFrame, detailed: pd.DataFrame, qc: pd.DataFrame, method: str, months: int, days: int, share: float, exclude_zero: bool, norm_references: dict[str, float]) -> bytes:
    how = "Excel supported oil" if method == "Excel: Oil Supported" else ("Excel support fraction × produced oil" if method == "Excel: Support Fraction" else f"Produced oil × {share:g}%")
    guide = pd.DataFrame([
        ("Scope", f"Uploaded period: {months} months. Raw data are total volumes in stb."),
        ("Oil Supported", how),
        ("Exclude zero-injection supported proxy", f"{exclude_zero}; applies to Produced Oil proxy methods only"),
        ("Raw Value of Water", "Sum(Oil Supported, stb) / Sum(Injected Water, stb), can exceed 1."),
        ("Normalized VoW (0–1)", "Raw group VoW / maximum raw VoW WITHIN THAT GROUPING LEVEL. Pattern, CPS and Manifold have separate references. Recomputed when filters change."),
        ("Pattern normalization reference (stb/stb)", f"{norm_references['Pattern']:.8f}"),
        ("CPS normalization reference (stb/stb)", f"{norm_references['CPS']:.8f}"),
        ("Manifold normalization reference (stb/stb)", f"{norm_references['Manifold']:.8f}"),
        ("Cross-plot breakdowns", "If a color option changes the actual point granularity (e.g. CPS x Reservoir), its plotted ratios are normalized by the maximum among THOSE displayed points. The displayed group is stated on the chart."),
        ("Normalization warning", "0–1 VoW is a relative score, NOT stb/stb and not the physically observed injection-to-oil ratio."),
        ("Aggregate VoW", "Ratio of sums, never an average of pattern VoW ratios."),
        ("Zero injection", "VoW is blank / undefined, not zero; retained in QC."),
        ("Interpretation", "Produced oil proxy is NOT demonstrated incremental oil caused by water injection."),
        ("CPS", "CPS number extracted from Pattern name, e.g. A-CPS2..., CPS-2-N."),
        ("CPS typo", "CSP4 variant mapped to CPS4; check original record."),
        ("Manifold data", "Uses the uploaded Manifold column directly; optional mapping only if missing/incomplete."),
        ("Manifold aggregation", "Group by unique Manifold ID (e.g. CPS-2/FM09); physical Manifold CPS takes priority over mismatched CPS in Pattern name. Mismatches flagged in QC."),
        ("Daily rate (mbd)", f"Full-period stb / {days} days / 1000; mbd means thousand stb/day."),
        ("VoW > 2", "Possible when numerator exceeds twice denominator; not a recovery factor. Confirm volumes, allocations and reporting periods."),
        ("VRR", "Reported source VRR is informational. Group averages are not field-corrected VRR."),
    ], columns=["Topic", "Rule / assumption"])
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        patterns.to_excel(writer, sheet_name="Pattern_Ranking", index=False)
        cps.to_excel(writer, sheet_name="CPS_Ranking", index=False)
        manifold.to_excel(writer, sheet_name="Manifold_Ranking", index=False)
        detailed.to_excel(writer, sheet_name="Calculated_Data", index=False)
        qc.to_excel(writer, sheet_name="QC", index=False)
        guide.to_excel(writer, sheet_name="Guideline", index=False)
        for sheet in writer.sheets.values():
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 25
            for cell in sheet[1]:
                from copy import copy
                font_style = copy(cell.font)
                font_style.bold = True
                font_style.color = "FFFFFF"
                cell.font = font_style
                from openpyxl.styles import PatternFill
                cell.fill = PatternFill(fill_type="solid", fgColor="173252")
            for col in sheet.columns:
                col_cells = list(col)
                max_len = max(len(str(c.value or "")) for c in col_cells[:250])
                from openpyxl.utils import get_column_letter
                sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = min(48, max(13, max_len + 2))
    return buffer.getvalue()


st.markdown("""
<style>
.block-container {padding-top: 1.65rem; padding-bottom: 2.5rem; max-width: 1550px;}
h1 {letter-spacing: -0.035em;}
[data-testid="stMetric"] {background: rgba(125,153,182,.07); border: 1px solid rgba(125,153,182,.17); border-radius: 12px; padding: 12px 16px;}
[data-testid="stMetricLabel"] {font-weight: 600;}
.stTabs [data-baseweb="tab"] {font-weight: 600;}
.small-note {font-size: .88rem; opacity: .85;}
</style>
""", unsafe_allow_html=True)

st.title("Pattern Waterflood Analytics")
st.caption("Rumaila | CPS / Pattern surveillance · Value of Water · Oil Supported · interactive ranking")

file = st.file_uploader("Browse your Pattern Excel or CSV file", type=["xlsx", "xls", "csv"], help="Upload Pattern_Data_Oct2026.xlsx or a file with equivalent columns.")
if file is None:
    st.info("Upload the October Pattern workbook to activate the dashboard. No source field data is stored in the app or uploaded to GitHub.")
    with st.expander("Required Excel structure", expanded=True):
        st.write("Required: **Pattern Name, Produced Oil (stb), Produced Water (stb), Injected Water (stb)**. Optional: Reservoir, Manifold, VRR, pressure, watercut, targets, and Oil Supported (stb).")
        st.write("Manifold analysis **automatically uses the Manifold column** in your uploaded Excel, when present. The optional mapping is only for incomplete files.")
        st.write("Automatic CPS parsing handles `A-CPS2-i9-R387`, `CPS-1-N` and flags the `CSP4` spelling variant.")
    st.stop()

ext = file.name.lower().rsplit(".", 1)[-1]
file_bytes = file.getvalue()
try:
    names = sheet_names(file_bytes, ext)
    selected_sheet = st.selectbox("Excel worksheet", names, index=0) if len(names) > 1 else names[0]
    raw = read_sheet(file_bytes, ext, selected_sheet)
    df, missing, matched_headers = standardise(raw)
except Exception as exc:
    st.error(f"Could not load the file: {exc}")
    st.stop()
if missing:
    st.error("Required columns not found: " + ", ".join(missing))
    st.write("Uploaded headers:", list(raw.columns))
    st.stop()
if df.empty:
    st.warning("No valid Pattern records detected.")
    st.stop()

# Make the source classification visible so a bad header cannot hide other
# reservoirs. All reservoirs are selected by default in the sidebar filters.
reservoir_header = matched_headers.get("Reservoir")
if reservoir_header is None:
    st.warning("Reservoir column not found. Patterns are marked 'Unspecified', not assumed to be Mishrif. Check the Excel headers.")
elif keyify(reservoir_header) not in {keyify(x) for x in FIELDS["Reservoir"]}:
    st.info(f"Reservoir detected automatically from column '{reservoir_header}'. "
            "Tip: rename its Excel header to 'Reservoir' for consistency.")
res_counts = df["Reservoir"].value_counts(dropna=False)
st.caption("Imported reservoir coverage: " + "  |  ".join(f"{name}: {count:,}" for name, count in res_counts.items()))

# -----------------------------------------------------------------------------
# Source Manifold column is preferred; mapping is optional only for missing labels.
# -----------------------------------------------------------------------------
source_manifold = matched_headers.get("Manifold")
source_mapped = int(df["Manifold"].ne("Unmapped").sum())
if source_manifold:
    st.success(f"Manifold source detected: **{source_manifold}** — {source_mapped:,} of {len(df):,} Pattern rows assigned.")
    st.caption("Manifold aggregation uses the Excel Manifold ID itself (e.g. CPS-2/FM09). Mismatches between Pattern CPS and Manifold CPS are flagged in QC rather than splitting one Manifold into multiple groups.")
else:
    st.warning("No Manifold column detected in this worksheet. Check selected sheet or supply a mapping below.")
with st.expander("Optional Manifold mapping / edits", expanded=False):
    st.caption("The Excel Manifold column is read automatically. Use the controls here only to fill or correct Manifold assignments; rows without a valid Manifold are excluded from Manifold plots.")
    mapping_template = (df[["Pattern", "CPS", "Reservoir", "Manifold"]]
                        .drop_duplicates(subset=["Pattern", "Reservoir"])
                        .sort_values(["CPS", "Reservoir", "Pattern"]).reset_index(drop=True))
    st.download_button("Download Pattern → Manifold mapping template (CSV)",
                       mapping_template.to_csv(index=False).encode("utf-8-sig"),
                       file_name="Pattern_Manifold_Mapping_Template.csv", mime="text/csv")
    mapping_file = st.file_uploader("Upload completed mapping CSV / Excel", type=["csv", "xlsx", "xls"],
                                    key="manifold_mapping_upload")
    if mapping_file is not None:
        try:
            mapping_bytes = mapping_file.getvalue()
            ext_mapping = mapping_file.name.rsplit(".", 1)[-1].lower()
            map_df = read_sheet(mapping_bytes, ext_mapping, sheet_names(mapping_bytes, ext_mapping)[0])
            df, mapping_notes = map_manifolds(df, map_df)
            for note in mapping_notes:
                st.caption(note)
            if any(x.startswith("Mapping requires") for x in mapping_notes):
                st.error("Mapping not applied. Please correct its headers.")
        except Exception as exc:
            st.error(f"Cannot read Manifold mapping: {exc}")
    manual_map = st.checkbox("Edit Manifold assignments directly in the app", False,
                             key="manifold_manual_toggle")
    if manual_map:
        st.caption("Fill Manifold values for the desired Patterns. Existing values are kept. For multiple Manifolds per Pattern, first supply a reliable allocated data table; do not assign volumes twice.")
        edit_df = (df[["Pattern", "CPS", "Reservoir", "Manifold"]]
                   .drop_duplicates(subset=["Pattern", "Reservoir"]).reset_index(drop=True))
        edited = st.data_editor(edit_df, hide_index=True, use_container_width=True, height=360,
                                disabled=["Pattern", "CPS", "Reservoir"], num_rows="fixed",
                                key="manifold_editor")
        df, mapping_notes = map_manifolds(df, edited)
        for note in mapping_notes:
            st.caption(note)
    assigned_count = int(df["Manifold"].ne("Unmapped").sum())
    st.caption(f"Assigned source rows: **{assigned_count:,} / {len(df):,}**. A valid manifold cross-plot needs source records with nonzero injection.")
# Reflect optional mapping edits in QC and source Manifold CPS checks too.
df["Manifold"] = df["Manifold"].map(clean_manifold)
df["Manifold CPS"] = df["Manifold"].map(manifold_cps)
df["Manifold CPS mismatch"] = (
    df["Manifold CPS"].notna() & df["CPS"].ne("Unassigned") & df["Manifold CPS"].ne(df["CPS"])
)

# -----------------------------------------------------------------------------
# Sidebar controls: global charts / period / support definition.
# -----------------------------------------------------------------------------
with st.sidebar:
    st.header("Data & calculation")
    n_months = st.number_input("Months represented by workbook", 1, 36, 3, 1)
    n_days = st.number_input("Days represented by workbook", min_value=1, max_value=3660,
                             value=int(n_months)*30, step=1,
                             help="Used to calculate daily rates: period total stb / actual number of days. Default is months × 30; set actual calendar days (e.g. 92) where known.")
    period = st.radio("Display volume as", [f"{n_months}-month total", "Average per month", "Average per day"], index=0,
                      help="Monthly = total volume / months; daily = total volume / selected days. VoW remains the ratio of unscaled volumes.")
    monthly = period == "Average per month"
    if n_months != 3:
        st.caption(f"Workbook is treated as a {n_months}-month period; this does not change the underlying totals.")
    unit = st.selectbox("Volume unit", list(UNIT_OPTIONS), index=1)
    factor, unit_label = UNIT_OPTIONS[unit]
    period_label = f"{unit_label}/day" if period == "Average per day" else (f"{unit_label}/month" if monthly else f"{unit_label} / period")
    if unit == "Mstb" and period == "Average per day":
        period_label = "mbd (thousand stb/day)"
    st.caption("**mbd** = thousand stock-tank barrels per day (Mstb/d). Enter exact period days for accurate daily rates.")
    support_methods = ["Produced Oil (screening proxy)", "Produced Oil × manual support %"]
    if df["Actual Supported Oil"].notna().any():
        support_methods.insert(0, "Excel: Oil Supported")
    if df["Support Fraction"].notna().any():
        support_methods.insert(0, "Excel: Support Fraction")
    method = st.selectbox("Oil Supported definition", support_methods,
                          index=support_methods.index("Produced Oil (screening proxy)"),
                          help="Actual waterflood-supported / incremental oil is not present in the October upload; Produced Oil is only a screening proxy.")
    share = 100.0
    if method == "Produced Oil × manual support %":
        share = st.slider("Oil attributable to water support (%)", 0, 100, 50)
    elif method == "Produced Oil (screening proxy)":
        share = 100.0
    exclude_zero_supported = st.checkbox("Exclude zero-injection Pattern oil from support proxy", True,
        help="When using an observed-oil proxy, a Pattern with zero injection cannot demonstrate contemporaneous injection support. Its production remains in Produced Oil, but is excluded from the Oil Supported proxy. Does not override explicitly uploaded actual supported oil.")
    if method.startswith("Produced Oil"):
        st.caption("⚠️ This is a screening assumption, not a measured incremental waterflood response.")
    # Display a relative 0–1 VoW score, preserving raw VoW in source data and tooltips.
    x_field = "VoW Normalized"
    x_caption = "Value of Water (normalised 0–1)"
    st.divider()
    st.header("Appearance")
    font = st.selectbox("Font", ["Arial", "Aptos", "Calibri", "Inter", "Open Sans", "Verdana"], index=0)
    title_size = st.slider("Chart title font", 12, 32, 19)
    axis_size = st.slider("Axis title font", 10, 26, 14)
    tick_size = st.slider("Axis tick font", 8, 20, 11)
    legend_size = st.slider("Legend font", 8, 22, 11)
    plot_height = st.slider("Chart height", 360, 1000, 570, 10)
    line_width = st.slider("Line thickness", 0.5, 9.0, 2.5, 0.5)
    marker_size = st.slider("Scatter marker / bubble max", 5, 42, 16)
    opacity = st.slider("Marker opacity", 0.2, 1.0, 0.83, 0.05)
    bg_preset = st.selectbox("Background style", ["White", "Soft grey", "Dark", "Custom"], index=0)
    if bg_preset == "Dark":
        paper, background, fg, grid_color = "#131E2D", "#18283B", "#EFF5FF", "#32465D"
    elif bg_preset == "Soft grey":
        paper, background, fg, grid_color = "#F5F8FC", "#F8FAFD", "#263849", "#DFE7EF"
    elif bg_preset == "Custom":
        paper = st.color_picker("Outer background", "#FFFFFF")
        background = st.color_picker("Plot area background", "#FFFFFF")
        fg = st.color_picker("Text color", "#1A2B41")
        grid_color = st.color_picker("Grid color", "#E1E8F0")
    else:
        paper, background, fg, grid_color = "#FFFFFF", "#FFFFFF", "#182536", "#E7EDF5"
    grid_on = st.checkbox("Show grid lines", True)
    legend_pos = st.selectbox("Legend position", ["Bottom", "Top", "Right", "Custom", "Hidden"])
    legend_x, legend_y, anchor = 0.5, -0.27, "center"
    if legend_pos == "Custom":
        legend_x = st.slider("Legend X", -0.5, 1.5, 0.5, 0.05)
        legend_y = st.slider("Legend Y", -0.8, 1.8, -0.25, 0.05)
        anchor = st.selectbox("Legend horizontal anchor", ["center", "left", "right"])
    with st.expander("Series & CPS colours"):
        oil_color = st.color_picker("Produced Oil", "#0C846F")
        water_color = st.color_picker("Produced Water", "#3695C4")
        injection_color = st.color_picker("Water Injection", "#D8832F")
        supported_color = st.color_picker("Oil Supported", "#1C9C85")
        cps_colors = {c: st.color_picker(c, color, key=f"clr_{c}") for c, color in DEFAULT_CPS_COLORS.items()}
    st.divider()
    st.header("Filters")
    cps_all = sorted(df["CPS"].unique().tolist(), key=cps_sort)
    selected_cps = st.multiselect("CPS", cps_all, default=cps_all)
    reservoirs = sorted(df["Reservoir"].unique().tolist())
    selected_res = st.multiselect("Reservoir", reservoirs, default=reservoirs)
    pattern_search = st.text_input("Search Pattern name", value="", placeholder="e.g. R387 or CPS-4")

cfg = Design(font=font, text=fg, paper=paper, plot=background, grid=grid_color,
             title_size=title_size, axis_size=axis_size, tick_size=tick_size,
             legend_size=legend_size, line_width=line_width, marker_size=marker_size,
             opacity=opacity, height=plot_height, legend_pos=legend_pos,
             legend_x=legend_x, legend_y=legend_y, legend_anchor=anchor,
             grid_on=grid_on)
filtered = df.loc[df["CPS"].isin(selected_cps) & df["Reservoir"].isin(selected_res)].copy()
if pattern_search:
    filtered = filtered.loc[filtered["Pattern"].str.contains(re.escape(pattern_search), case=False, na=False)]
if filtered.empty:
    st.warning("Current filters return no Pattern records. Adjust the sidebar selections.")
    st.stop()

work = support_scenario(filtered, method, share)
if exclude_zero_supported and method.startswith("Produced Oil"):
    work.loc[work["Injection"].isna() | work["Injection"].le(0), "Supported"] = 0.0
pat = attach_display(rollup(work, ["Pattern", "CPS", "Reservoir", "Manifold"]), factor, n_months, period, n_days)
cps = attach_display(rollup(work, ["CPS"]), factor, n_months, period, n_days)
cps = cps.sort_values("CPS", key=lambda x: x.map(cps_sort)).reset_index(drop=True)
mapped_work = work.loc[work["Manifold"].ne("Unmapped")].copy()
# A Manifold identifier normally already includes its CPS (CPS-2/FM09). Group
# using that identifier; do not split it based on CPS parsed from Pattern names.
manifold = attach_display(manifold_rollup(mapped_work), factor, n_months, period, n_days) if not mapped_work.empty else pd.DataFrame()
# Separate references: each aggregation level is independently scaled to 0–1.
# Never let a high Pattern raw ratio force all CPS/Manifold scores below 1.
by_cps_res = rollup(work, ["CPS", "Reservoir"])
by_man_res = manifold_rollup(mapped_work, ["Reservoir"]) if not mapped_work.empty else pd.DataFrame()
norm_references = {
    "Pattern": vow_normalization_reference(pat),
    "CPS": vow_normalization_reference(cps),
    "Manifold": vow_normalization_reference(manifold),
}
pat = normalise_vow(pat, norm_references["Pattern"])
cps = normalise_vow(cps, norm_references["CPS"])
if not manifold.empty:
    manifold = normalise_vow(manifold, norm_references["Manifold"])
valid_manifold = manifold.loc[manifold["VoW Valid"] & manifold["VoW Technical"].notna()].copy() if not manifold.empty else pd.DataFrame()
valid_pat = pat[pat["VoW Valid"] & pat["VoW Technical"].notna()].copy()
valid_cps = cps[cps["VoW Valid"] & cps["VoW Technical"].notna()].copy()
total_oil = float(work["Oil"].sum())
total_water = float(work["Water"].sum())
total_inj = float(work["Injection"].sum())
total_supported = float(work["Supported"].sum())
weighted_vow = total_supported / total_inj if total_inj > 0 else np.nan

# QC over ORIGINAL rows, scoped to selected CPS/reservoir/pattern.
qc_frames = []
for condition, text_label in [
    (filtered["Injection"].eq(0), "Zero water injection — VoW undefined"),
    (filtered["Manifold"].eq("Unmapped"), "Manifold assignment missing — excluded from Manifold aggregation"),
    (filtered["Manifold CPS mismatch"], "Pattern CPS and Manifold CPS differ — Manifold value used for Manifold rollup; verify assignment"),
    (work["Injection"].gt(0) & work["Supported"].gt(2 * work["Injection"]), "VoW > 2 — review context; mathematically possible"),
    (filtered["Injection"].isna(), "Missing injected volume"),
    (filtered["Oil"].isna(), "Missing produced oil volume"),
    (filtered["Water"].isna(), "Missing produced water volume"),
    (filtered["CPS QC"].ne(""), "CPS name parse / spelling issue"),
    (filtered["VRR"].gt(10), "High source VRR (>10) — check"),
    (filtered["Injection"].lt(0) | filtered["Oil"].lt(0) | filtered["Water"].lt(0), "Negative source volume"),
]:
    if condition.any():
        a = filtered.loc[condition, ["Source Row", "Pattern", "CPS", "Reservoir", "Manifold", "Oil", "Water", "Injection", "VRR", "CPS QC"]].copy()
        a["QC Reason"] = text_label
        qc_frames.append(a)
qc_table = pd.concat(qc_frames, ignore_index=True) if qc_frames else pd.DataFrame(columns=["Source Row", "Pattern", "CPS", "Reservoir", "Manifold", "QC Reason"])

# Top summary, intentionally independent of tab selection to avoid empty overview after navigation.
st.caption(f"{len(filtered):,} source rows · {len(pat):,} Patterns · {len(cps):,} CPS groups · {len(valid_manifold):,} valid Manifold groups · {n_days} period days · Manifold IDs from source Excel")
met1, met2, met3, met4, met5 = st.columns(5)
met1.metric(f"Produced Oil ({period_label})", pretty(float(display_volume(pd.Series([total_oil]), factor, n_months, period, n_days).iloc[0]), 1))
met2.metric(f"Produced Water ({period_label})", pretty(float(display_volume(pd.Series([total_water]), factor, n_months, period, n_days).iloc[0]), 1))
met3.metric(f"Injected Water ({period_label})", pretty(float(display_volume(pd.Series([total_inj]), factor, n_months, period, n_days).iloc[0]), 1))
met4.metric(f"Oil Supported ({period_label})", pretty(float(display_volume(pd.Series([total_supported]), factor, n_months, period, n_days).iloc[0]), 1))
met5.metric("Weighted VoW (raw stb/stb)", pretty(weighted_vow, 3))
st.caption(
    f"Level-specific VoW references (raw stb/stb): **Pattern {norm_references['Pattern']:.3f}** · "
    f"**CPS {norm_references['CPS']:.3f}** · **Manifold {norm_references['Manifold']:.3f}**. "
    "For each chart, the highest valid VoW within the plotted grouping becomes 1.000; "
    "raw weighted portfolio VoW above is deliberately kept in original stb/stb units."
)
if method.startswith("Produced Oil"):
    st.warning("**Interpretation:** Oil Supported is an assumption based on Produced Oil, not measured incremental oil attributable to injection. Use an actual supported-oil column for defensible causal Value of Water.", icon="⚠️")

# -----------------------------------------------------------------------------
# Plot construction and tabs.
# -----------------------------------------------------------------------------
tabs = st.tabs(["Overview", "VoW Cross-plot", "Pattern Ranking", "CPS Ranking", "Manifold Ranking", "Surveillance", "Data & Guideline"])

with tabs[0]:
    st.subheader("CPS production & injection overview")
    grouped = go.Figure()
    for fld, name, colour in [
        ("Oil Display", "Produced Oil", oil_color), ("Water Display", "Produced Water", water_color),
        ("Injection Display", "Injected Water", injection_color),
    ]:
        grouped.add_trace(go.Bar(x=cps["CPS"], y=cps[fld], name=name,
                                 marker_color=colour, opacity=cfg.opacity,
                                 hovertemplate=f"%{{x}}<br>{name}: %{{y:,.2f}} {period_label}<extra></extra>"))
    grouped.update_layout(barmode="group", bargap=0.22)
    apply_style(grouped, cfg, "Produced oil, produced water & injected water — by CPS", "CPS", f"Volume ({period_label})")
    plot(grouped, "overview_cps_volumes")

    left, right = st.columns(2)
    with left:
        wc = go.Figure()
        wc.add_trace(go.Bar(x=cps["CPS"], y=cps["Watercut"] * 100,
                            marker_color=[cps_colors.get(x, "#789") for x in cps["CPS"]],
                            text=(cps["Watercut"] * 100).map(lambda v: f"{v:.1f}%" if pd.notna(v) else ""), textposition="outside"))
        apply_style(wc, cfg, "Produced water cut — by CPS", "CPS", "Water cut (%)", show_legend=False, height=480)
        wc.update_yaxes(rangemode="tozero")
        plot(wc, "overview_watercut")
    with right:
        cw = go.Figure()
        cw.add_trace(go.Bar(x=cps["CPS"], y=cps["Supported Display"], name="Oil Supported", marker_color=supported_color))
        apply_style(cw, cfg, "Oil Supported — by CPS", "CPS", f"Oil Supported ({period_label})", show_legend=False, height=480)
        plot(cw, "overview_oil_supported")
    st.caption("Volumes are aggregated by CPS from the Pattern names. Water cut is recomputed from total produced water / (total oil + total water), not averaged across Patterns.")

with tabs[1]:
    st.subheader("Value of Water vs Oil Supported")
    a, b, c, d = st.columns(4)
    with a:
        plot_level = st.radio("Aggregation", ["Pattern", "CPS", "Manifold"], horizontal=True, key="cross_level")
    with b:
        color_by = st.selectbox("Color by", ["CPS", "Reservoir", "Pattern", "Manifold"], key="cross_color_by")
    with c:
        bubble_mode = st.selectbox("Marker sizing", ["Fixed", "Injected Water", "Oil Supported"], key="cross_bubble")
    with d:
        point_labels = st.selectbox("Point labels", ["Off", "Top 10", "All"], key="cross_labels")

    opts_a, opts_b, opts_c = st.columns([1, 1, 1])
    with opts_a:
        use_quadrants = st.checkbox("Show median-based quadrants", True)
    with opts_b:
        trendline = st.checkbox("Show linear trend line", False)
    with opts_c:
        show_category_legend = st.checkbox("Show category legend", value=color_by != "Pattern", key=f"cross_show_legend_{color_by}")

    oil_axis = st.radio("Cross-plot Oil Supported Y-axis", ["Average per day (mbd)", "Follow global volume setting"],
                        index=0, horizontal=True, key="cross_y_unit",
                        help="Default: supported oil (stb) / days in uploaded period / 1000. The X-axis VoW is unchanged.")
    y_caption = "Oil Supported (mbd)" if oil_axis.startswith("Average") else f"Oil Supported ({period_label})"
    # Stable colors independent of filtering. CPS retains the editable sidebar colors.
    palette_name = "Plotly"
    custom_color = {}
    if color_by != "CPS":
        palette_name = st.selectbox("Color palette", ["Plotly", "Safe", "Bold", "Dark24", "Alphabet", "Pastel", "Set3"], key="cross_palette")
        if color_by == "Reservoir":
            with st.expander("Reservoir colors (custom)"):
                for idx, res_name in enumerate(sorted(df["Reservoir"].unique().tolist())):
                    pcols = getattr(px.colors.qualitative, palette_name)
                    custom_color[res_name] = st.color_picker(res_name, pcols[idx % len(pcols)], key=f"cross_res_color_{idx}")

    custom_title = st.text_input("Cross-plot chart title", "Value of Water vs Oil Supported", key="cross_title")
    display_level = plot_level
    if plot_level == "Pattern":
        scatter_df = valid_pat.copy()
    elif plot_level == "Manifold":
        if mapped_work.empty:
            scatter_df = pd.DataFrame()
            st.info("No Manifold mapping has been provided. Open 'Manifold mapping' above, upload a Pattern → Manifold CSV or edit the table.")
        elif color_by == "Pattern":
            scatter_df = valid_pat.loc[valid_pat["Manifold"].ne("Unmapped")].copy()
            display_level = "Pattern breakdown within mapped Manifolds"
        elif color_by == "Reservoir":
            scatter_df = attach_display(manifold_rollup(mapped_work, ["Reservoir"]), factor, n_months, period, n_days)
            scatter_df = scatter_df.loc[scatter_df["VoW Valid"]].copy()
            display_level = "Manifold × Reservoir"
        else:
            scatter_df = valid_manifold.copy()
            display_level = "CPS × Manifold"
    elif color_by == "Manifold":
        scatter_df = valid_manifold.copy()
        display_level = "Manifold breakdown by CPS"
        if scatter_df.empty:
            st.info("Color by Manifold requires a Pattern → Manifold mapping.")
    elif color_by == "Reservoir":
        # One point per CPS / Reservoir, so no potentially mixed-reservoir CPS is
        # incorrectly painted as belonging to just one reservoir.
        scatter_df = rollup(work, ["CPS", "Reservoir"])
        scatter_df = attach_display(scatter_df, factor, n_months, period, n_days)
        scatter_df = scatter_df.loc[scatter_df["VoW Valid"]].copy()
        display_level = "CPS × Reservoir"
        st.caption("Reservoir coloring at CPS level shows one aggregated point per CPS–Reservoir combination, rather than assigning an arbitrary reservoir to the entire CPS.")
    elif color_by == "Pattern":
        # A CPS aggregate has no unique Pattern. Keep values at true Pattern
        # granularity, grouped by selected CPS, instead of assigning false colors.
        scatter_df = valid_pat.copy()
        display_level = "Pattern breakdown by CPS"
        st.info("Color by Pattern requires Pattern-level points. Displaying Pattern breakdown within the selected CPS groups; CPS totals remain available under Color by CPS.")
    else:
        scatter_df = valid_cps.copy()

    if scatter_df.empty:
        st.info("No records with a positive injected-water volume and a defined Oil Supported value can be plotted. See QC.")
    else:
        # Always re-scale to the CURRENTLY DISPLAYED point granularity.
        # Pattern / CPS / Manifold use their own maxima; a color breakdown
        # (CPS x Reservoir, Manifold x Reservoir, etc.) is re-scaled to its
        # own plotted group maximum so there is exactly one 1.000 point.
        scatter_norm_reference = vow_normalization_reference(scatter_df)
        scatter_df = normalise_vow(scatter_df, scatter_norm_reference)
        scatter_df["x"] = scatter_df[x_field]
        scatter_df["y"] = scatter_df["Supported mbd"] if oil_axis.startswith("Average") else scatter_df["Supported Display"]
        scatter_df = scatter_df.replace([np.inf, -np.inf], np.nan).dropna(subset=["x", "y"])
        if "Pattern" in scatter_df.columns:
            scatter_df["Point"] = scatter_df["Pattern"]
        elif "Manifold" in scatter_df.columns:
            scatter_df["Point"] = scatter_df["CPS"] + " · " + scatter_df["Manifold"]
            if "Reservoir" in scatter_df.columns:
                scatter_df["Point"] += " · " + scatter_df["Reservoir"]
        elif "Reservoir" in scatter_df.columns:
            scatter_df["Point"] = scatter_df["CPS"] + " · " + scatter_df["Reservoir"]
        else:
            scatter_df["Point"] = scatter_df["CPS"]
        scatter_df["Bubble"] = (scatter_df["Injection Display"] if bubble_mode == "Injected Water"
                                else scatter_df["Supported Display"] if bubble_mode == "Oil Supported"
                                else pd.Series([1] * len(scatter_df), index=scatter_df.index))
        scatter_df["Bubble"] = scatter_df["Bubble"].clip(lower=0).fillna(0)
        scatter_df["Water cut (%)"] = scatter_df["Watercut"] * 100
        scatter_df["Injected Water"] = scatter_df["Injection Display"]
        scatter_df["Oil Supported"] = scatter_df["y"]
        scatter_df["Oil Supported (mbd)"] = scatter_df["Supported mbd"]
        scatter_df["Raw VoW (stb/stb)"] = scatter_df["VoW Technical"]
        scatter_df["Normalised VoW (0–1)"] = scatter_df["VoW Normalized"]
        if scatter_df.empty:
            st.info("No valid numerical points after filtering.")
        else:
            if color_by == "CPS":
                color_map = cps_colors
            elif color_by == "Reservoir":
                color_map = custom_color
            elif color_by == "Manifold":
                pal = getattr(px.colors.qualitative, palette_name)
                color_map = {m: pal[i % len(pal)] for i, m in enumerate(sorted(df["Manifold"].unique().tolist()))}
            else:
                palette = getattr(px.colors.qualitative, palette_name)
                # Stable global ordering maintains the same color for a Pattern
                # after filtering, when a small number of Patterns is selected.
                patterns_all = sorted(df["Pattern"].dropna().astype(str).unique().tolist())
                color_map = {name: palette[i % len(palette)] for i, name in enumerate(patterns_all)}

            hover = {"CPS": True, "Injected Water": ":,.2f", "Oil Supported": ":,.2f",
                     "Oil Supported (mbd)": ":,.2f",
                     "Raw VoW (stb/stb)": ":.4f", "Normalised VoW (0–1)": ":.4f", "Water cut (%)": ":.1f",
                     "x": False, "y": False, "Bubble": False}
            if "Reservoir" in scatter_df.columns:
                hover["Reservoir"] = True
            if "Manifold" in scatter_df.columns:
                hover["Manifold"] = True
            if "Pattern" in scatter_df.columns:
                hover["Pattern"] = True

            fig_sc = px.scatter(
                scatter_df, x="x", y="y", color=color_by, size="Bubble", size_max=cfg.marker_size,
                hover_name="Point", color_discrete_map=color_map, hover_data=hover,
                category_orders={"CPS": sorted(df["CPS"].unique(), key=cps_sort)},
            )
            fig_sc.update_traces(marker=dict(opacity=cfg.opacity, line=dict(width=max(0.5, cfg.line_width / 2), color=cfg.paper)))
            if use_quadrants and len(scatter_df) > 1:
                mid_x, mid_y = float(scatter_df["x"].median()), float(scatter_df["y"].median())
                fig_sc.add_vline(x=mid_x, line_dash="dash", line_color="#8793A4", line_width=cfg.line_width)
                fig_sc.add_hline(y=mid_y, line_dash="dash", line_color="#8793A4", line_width=cfg.line_width)
                st.caption(f"Quadrants use {display_level.lower()} medians: normalised VoW = {mid_x:,.3f}; Oil Supported = {mid_y:,.2f} {'mbd' if oil_axis.startswith('Average') else period_label}.")
            if trendline and len(scatter_df) >= 3 and scatter_df["x"].nunique() > 1:
                slope, offset = np.polyfit(scatter_df["x"], scatter_df["y"], 1)
                xs = np.linspace(scatter_df["x"].min(), scatter_df["x"].max(), 100)
                fig_sc.add_trace(go.Scatter(x=xs, y=slope * xs + offset, mode="lines", name="OLS trend",
                                             line=dict(color="#D64E60", width=cfg.line_width + 1, dash="dot"), hoverinfo="skip"))
            if point_labels != "Off":
                to_label = scatter_df if point_labels == "All" else scatter_df.nlargest(min(10, len(scatter_df)), "y")
                fig_sc.add_trace(go.Scatter(x=to_label["x"], y=to_label["y"], mode="text", text=to_label["Point"],
                                             textposition="top center", textfont=dict(size=max(9, cfg.tick_size - 1), color=cfg.text),
                                             showlegend=False, hoverinfo="skip", cliponaxis=False))
            apply_style(fig_sc, cfg, f"{custom_title} — {display_level}", x_caption, y_caption, show_legend=show_category_legend)
            manual_scatter_axes = st.checkbox("Set manual X/Y axis ranges", False, key="manual_scatter_axes")
            if manual_scatter_axes:
                min_x, max_x = float(scatter_df["x"].min()), float(scatter_df["x"].max())
                min_y, max_y = float(scatter_df["y"].min()), float(scatter_df["y"].max())
                pad_x, pad_y = max((max_x - min_x) * 0.08, 0.01), max((max_y - min_y) * 0.08, 0.01)
                xx1, xx2, yy1, yy2 = st.columns(4)
                with xx1:
                    lo_x = st.number_input("X minimum", min_value=0.0, max_value=1.0, value=max(0.0,min_x-pad_x), format="%.4f", key="x_axis_min")
                with xx2:
                    hi_x = st.number_input("X maximum", min_value=0.0, max_value=1.0, value=min(1.0,max_x+pad_x), format="%.4f", key="x_axis_max")
                with yy1:
                    lo_y = st.number_input("Y minimum", value=min_y - pad_y, format="%.2f", key="y_axis_min")
                with yy2:
                    hi_y = st.number_input("Y maximum", value=max_y + pad_y, format="%.2f", key="y_axis_max")
                if lo_x < hi_x and lo_y < hi_y:
                    fig_sc.update_xaxes(range=[lo_x, hi_x])
                    fig_sc.update_yaxes(range=[lo_y, hi_y])
                else:
                    st.warning("Axis maximum must be greater than minimum.")
            if not manual_scatter_axes:
                fig_sc.update_xaxes(range=[0, 1], tick0=0, dtick=0.2)
            plot(fig_sc, f"vow_scatter_{plot_level.lower()}_{color_by.lower()}")
            st.download_button("Download cross-plot as interactive HTML", fig_sc.to_html(full_html=True, include_plotlyjs=True).encode("utf-8"),
                               file_name=f"VoW_vs_OilSupported_{plot_level}_ColorBy_{color_by}.html", mime="text/html", key="html_cross")
            # True weighted aggregation: sum(supported) / sum(injected). Not average of ratios.
            st.caption(
                f"VoW (0–1) for **{display_level}** = raw VoW ÷ **{scatter_norm_reference:.3f} stb/stb**, "
                "the highest raw ratio **among the points shown on this plot**. "
                "Changing Aggregation refreshes the reference. If Color by creates a "
                "CPS × Reservoir or other detailed breakdown, the plot shows and scales "
                "that actual breakdown (not unsplit CPS totals). "
                "Original ratios remain in hover and exports."
            )

with tabs[2]:
    st.subheader("Pattern-by-Pattern ranking")
    rank_measure = st.selectbox("Rank Patterns by", ["VoW Normalized", "Supported Display", "Injection Display", "Oil Display", "Watercut"], key="rank_measure")
    rank_max = max(1, min(len(pat), 200))
    num_to_show = st.slider("Number of Patterns in chart", 1 if rank_max < 5 else 5, rank_max, min(len(pat), 30), key="rank_n")
    minimum_inj = st.number_input("Minimum injected volume (stb / original full period)", min_value=0.0, value=0.0, step=5_000.0, key="rank_min_inj")
    ranking = pat.loc[pat["VoW Valid"] & pat["Injection"].ge(minimum_inj)].copy()
    ranking = ranking.sort_values(rank_measure, ascending=False, na_position="last").reset_index(drop=True)
    ranking.insert(0, "Rank", np.arange(1, len(ranking) + 1))
    ranked_n = ranking.head(num_to_show).sort_values(rank_measure, ascending=True)
    if not ranked_n.empty:
        chart = go.Figure(go.Bar(x=ranked_n[rank_measure], y=ranked_n["Pattern"], orientation="h",
                                 marker_color=[cps_colors.get(v, "#789") for v in ranked_n["CPS"]],
                                 customdata=ranked_n[["CPS", "Oil Display", "Injection Display"]].fillna(0).to_numpy(),
                                 hovertemplate="%{y}<br>Score: %{x:,.3f}<br>CPS: %{customdata[0]}<br>Oil: %{customdata[1]:,.1f}<br>Injection: %{customdata[2]:,.1f}<extra></extra>"))
        bar_height = max(470, min(1900, int(num_to_show * 28 + 180)))
        labels = {"VoW Normalized": "Value of Water (normalised 0–1)",
                  "Supported Display": f"Oil Supported ({period_label})", "Injection Display": f"Injection ({period_label})",
                  "Oil Display": f"Oil ({period_label})", "Watercut": "Produced-water fraction"}
        apply_style(chart, cfg, f"Top Patterns ranked by {rank_measure}", labels[rank_measure], "Pattern", show_legend=False, height=bar_height)
        chart.update_layout(margin=dict(l=155, r=70, t=100, b=95))
        if rank_measure == "VoW Normalized":
            chart.update_xaxes(range=[0, 1])
        plot(chart, "pattern_ranking_bar")
    else:
        st.info("No Patterns satisfy the chosen injected-volume threshold.")
    table = ranking[["Rank", "Pattern", "CPS", "Reservoir", "Manifold", "VoW Normalized", "VoW Technical",
                     "Supported Display", "Oil Display", "Water Display", "Injection Display", "Watercut", "VRR", "Res Pressure", "FBHP Trend"]].copy()
    table = table.rename(columns={"Supported Display": f"Supported ({period_label})", "Oil Display": f"Oil ({period_label})",
                                  "Water Display": f"Produced Water ({period_label})", "Injection Display": f"Injected Water ({period_label})",
                                  "Watercut": "Watercut (fraction)", "Res Pressure": "Reservoir Pressure (psi)"})
    st.dataframe(table, use_container_width=True, hide_index=True, height=445)
    st.caption("Ranking includes only Patterns with positive injection and a defined supported-oil value. Zero-injection patterns are retained in the QC and raw data exports.")

with tabs[3]:
    st.subheader("CPS ranking & allocation")
    cps_rank_measure = st.selectbox("CPS score", ["VoW Normalized", "Supported Display", "Injection Display"], key="cps_rank_metric")
    rank_cps = valid_cps.sort_values(cps_rank_measure, ascending=False).reset_index(drop=True)
    if not rank_cps.empty:
        rank_cps.insert(0, "Rank", np.arange(1, len(rank_cps) + 1))
        order = rank_cps.sort_values(cps_rank_measure)
        cps_fig = go.Figure(go.Bar(x=order[cps_rank_measure], y=order["CPS"], orientation="h",
                                   marker_color=[cps_colors.get(i, "#789") for i in order["CPS"]],
                                   text=order[cps_rank_measure].map(lambda v: f"{v:,.3f}" if pd.notna(v) else ""), textposition="outside"))
        apply_style(cps_fig, cfg, "CPS ranking — weighted by injected water", cps_rank_measure, "CPS", show_legend=False, height=460)
        if cps_rank_measure == "VoW Normalized":
            cps_fig.update_xaxes(range=[0, 1])
        plot(cps_fig, "cps_ranking")
        cpstable = rank_cps[["Rank", "CPS", "Patterns", "VoW Normalized", "VoW Technical", "Supported Display",
                             "Oil Display", "Water Display", "Injection Display", "Watercut", "VRR"]].copy()
        cpstable = cpstable.rename(columns={"Supported Display": f"Supported ({period_label})", "Oil Display": f"Oil ({period_label})",
                                             "Water Display": f"Water ({period_label})", "Injection Display": f"Injection ({period_label})",
                                             "VRR": "Mean reported VRR (not aggregated VRR)"})
        st.dataframe(cpstable, use_container_width=True, hide_index=True)
    else:
        st.info("No CPS with positive injection and defined support.")
    st.caption("CPS raw VoW = total supported oil / total injected water. Ranking displays its 0–1 normalised score. Do not use a simple average of Pattern efficiencies.")

with tabs[4]:
    st.subheader("Manifold-by-Manifold ranking")
    if manifold.empty:
        st.info("No Manifold data are assigned. Use the Manifold mapping control above to import an approved Pattern → Manifold table or fill it manually.")
    else:
        st.caption(f"Mapped Pattern rows: {len(mapped_work):,} of {len(work):,}. **Unmapped rows are excluded only from Manifold analysis**, not CPS / Pattern calculations.")
        rankman = valid_manifold.sort_values("VoW Normalized", ascending=False).reset_index(drop=True)
        if not rankman.empty:
            plotman = rankman.sort_values("VoW Normalized")
            figman = go.Figure(go.Bar(x=plotman["VoW Normalized"], y=plotman["Manifold"],
                                     orientation="h", marker_color=[cps_colors.get(c, "#789") for c in plotman["CPS"]],
                                     text=plotman["VoW Normalized"].map(lambda v: f"{v:.3f}"), textposition="outside"))
            apply_style(figman, cfg, "Manifold ranking — weighted by injected water", "Value of Water (normalised 0–1)",
                        "Manifold", show_legend=False, height=max(480, min(1600, len(plotman)*33 + 180)))
            figman.update_xaxes(range=[0,1])
            plot(figman, "manifold_ranking")
            st.dataframe(rankman[["CPS", "Manifold", "Patterns", "VoW Normalized", "VoW Technical", "Supported mbd", "Injection mbd",
                                    "Supported Display", "Injection Display", "Watercut"]],
                         use_container_width=True, hide_index=True)
        else:
            st.warning("Manifold assignments were found, but no mapped Manifolds have valid positive injection and supported-oil volumes.")

with tabs[5]:
    st.subheader("Surveillance & diagnostic plots")
    obs1, obs2 = st.columns(2)
    with obs1:
        diag1 = px.scatter(valid_pat, x="VRR", y="Watercut", color="CPS", color_discrete_map=cps_colors,
                           hover_name="Pattern", size="Injection Display", size_max=cfg.marker_size)
        diag1.update_traces(marker=dict(opacity=cfg.opacity, line=dict(width=cfg.line_width / 2, color=cfg.paper)))
        apply_style(diag1, cfg, "Reported VRR vs Produced Watercut", "Reported Pattern VRR", "Produced watercut (fraction)", height=490)
        plot(diag1, "diag_vrr_watercut")
    with obs2:
        diag2 = px.scatter(valid_pat, x="VoW Normalized", y="Watercut", color="CPS", color_discrete_map=cps_colors,
                           hover_name="Pattern", size="Injection Display", size_max=cfg.marker_size)
        diag2.update_traces(marker=dict(opacity=cfg.opacity, line=dict(width=cfg.line_width / 2, color=cfg.paper)))
        apply_style(diag2, cfg, "Watercut vs Value of Water", "VoW (normalised 0–1)", "Produced watercut (fraction)", height=490)
        diag2.update_xaxes(range=[0,1])
        plot(diag2, "diag_vow_watercut")
    obs3, obs4 = st.columns(2)
    with obs3:
        pressure = pat.dropna(subset=["Res Pressure", "FBHP Trend"])
        if len(pressure):
            pres_fig = px.scatter(pressure, x="Res Pressure", y="FBHP Trend", color="CPS", color_discrete_map=cps_colors,
                                  hover_name="Pattern", size="Injection Display", size_max=cfg.marker_size)
            apply_style(pres_fig, cfg, "Reservoir pressure vs FBHP trend", "Average reservoir pressure (psi)", "Average FBHP trend (psi/day)", height=490)
            plot(pres_fig, "diag_pressure_fbhp")
        else:
            st.info("Reservoir pressure / FBHP trend data missing for selected Patterns.")
    with obs4:
        if not valid_pat.empty:
            hist = px.histogram(valid_pat, x="VoW Normalized", color="CPS", color_discrete_map=cps_colors,
                                nbins=min(40, max(8, len(valid_pat) // 4)), barmode="overlay", opacity=cfg.opacity)
            apply_style(hist, cfg, "Distribution of Pattern Value of Water", "VoW (normalised 0–1)", "Pattern count", height=490)
            hist.update_xaxes(range=[0,1])
            plot(hist, "diag_vow_histogram")
        else:
            st.info("No valid Patterns for VoW distribution.")
    st.divider()
    st.markdown("**Custom cross-plot lab**")
    available = {
        "Value of Water (normalised)": ("VoW Normalized", "0–1"),
        "Produced Oil": ("Oil Display", period_label),
        "Oil Supported": ("Supported Display", period_label),
        "Produced Water": ("Water Display", period_label),
        "Water Injection": ("Injection Display", period_label),
        "Watercut": ("Watercut", "fraction"),
        "Reported VRR": ("VRR", "source value"),
        "Reservoir Pressure": ("Res Pressure", "psi"),
        "Delta P to Psat": ("Delta P", "psi"),
        "FBHP Trend": ("FBHP Trend", "psi/day"),
        "Agreed Target": ("Agreed Target", "stb source period"),
        "RMT Target": ("RMT Target", "stb source period"),
    }
    cx, cy, clev = st.columns([1, 1, 1])
    with cx:
        x_custom = st.selectbox("X variable", list(available), index=4)
    with cy:
        y_custom = st.selectbox("Y variable", list(available), index=2)
    with clev:
        custom_level = st.selectbox("Grouping", ["Pattern", "CPS"], key="custom_level")
    lab_data = (pat if custom_level == "Pattern" else cps).copy()
    xv, xu = available[x_custom]
    yv, yu = available[y_custom]
    lab_data = lab_data.replace([np.inf, -np.inf], np.nan).dropna(subset=[xv, yv])
    if len(lab_data):
        lab_fig = px.scatter(lab_data, x=xv, y=yv, color="CPS", color_discrete_map=cps_colors,
                              hover_name="Pattern" if custom_level == "Pattern" else "CPS")
        lab_fig.update_traces(marker=dict(size=cfg.marker_size, opacity=cfg.opacity,
                                          line=dict(width=cfg.line_width / 2, color=cfg.paper)))
        apply_style(lab_fig, cfg, f"{y_custom} vs {x_custom} — {custom_level}", f"{x_custom} ({xu})", f"{y_custom} ({yu})")
        if xv == "VoW Normalized":
            lab_fig.update_xaxes(range=[0,1])
        if yv == "VoW Normalized":
            lab_fig.update_yaxes(range=[0,1])
        plot(lab_fig, "custom_scatter_lab")
    else:
        st.info("These X and Y variables have no overlapping non-empty observations.")
    st.caption("Pressure data are missing for some Patterns; these are dropped from pressure plots only. Supplied VRR is not recalculated from stock-tank volumes.")

with tabs[6]:
    st.subheader("Source data, QC & calculation guideline")
    qc_col1, qc_col2, qc_col3, qc_col4 = st.columns(4)
    qc_col1.metric("Original selected records", len(filtered))
    qc_col2.metric("Zero-injection records", int(filtered["Injection"].eq(0).sum()))
    qc_col3.metric("Unparsed / corrected CPS names", int(filtered["CPS QC"].ne("").sum()))
    qc_col4.metric("Manifold unmapped rows", int(filtered["Manifold"].eq("Unmapped").sum()))
    st.caption(f"Manifold CPS mismatch flags: {int(filtered['Manifold CPS mismatch'].sum())}. Rows stay assigned to their explicit Manifold; inspect the QC exceptions if a Pattern's CPS differs.")
    if not qc_table.empty:
        with st.expander(f"QC exceptions ({len(qc_table)} flags)", expanded=False):
            st.dataframe(qc_table, use_container_width=True, hide_index=True, height=300)
    with st.expander("Column matching", expanded=False):
        st.dataframe(pd.DataFrame([{"Dashboard field": a, "Excel source header": b} for a, b in matched_headers.items()]), hide_index=True)
    with st.expander("Source plus calculated columns", expanded=False):
        st.dataframe(work, use_container_width=True, hide_index=True, height=340)
    st.markdown("### Guideline — Value of Water")
    st.markdown("**1. Identify Oil Supported.** When an actual injection-attributable oil estimate is available, it should be used. The October Excel contains Produced Oil, but no direct Oil Supported / incremental oil column.")
    st.latex(r"Q_{o,\mathrm{supported}}=Q_{o,\mathrm{produced}}\times f_{\mathrm{support}}")
    st.markdown(f"Current method: **{method}**. Current oil-support percentage (when relevant): **{share:g}%**. Exclude zero-injection Pattern oil from produced-oil proxy: **{exclude_zero_supported}** (only applies to Produced Oil methods).")
    st.markdown("**2. Calculate original Value of Water (VoW).**")
    st.latex(r"\mathrm{VoW}_{\mathrm{technical}}=\frac{\sum Q_{o,\mathrm{supported}}\ [\mathrm{stb}]}{\sum W_{\mathrm{inj}}\ [\mathrm{stb}]}\ \ [\mathrm{stb\ oil/stb\ injected\ water}]")
    st.markdown("This is the original volume ratio (stb/stb). It can exceed 1 and is undefined for zero injection. It is not proof of incremental recovery.")
    st.markdown("**Display normalisation (0–1)** — separate max for each aggregation level:")
    st.latex(r"\mathrm{VoW}_{\mathrm{norm},L}=\frac{\mathrm{VoW}_{\mathrm{raw},L}}{\max_{g\in L}(\mathrm{VoW}_{\mathrm{raw},g})},\quad L\in\{\mathrm{Pattern,\ CPS,\ Manifold}\}")
    st.markdown(
        f"Current raw maxima (**selected filters**): Pattern **{norm_references['Pattern']:.4f}**, "
        f"CPS **{norm_references['CPS']:.4f}**, "
        f"Manifold **{norm_references['Manifold']:.4f}** stb/stb. "
        "Each level is normalised independently, so its highest valid group scores 1.000. "
        "For example, if max Pattern VoW is 4.32, that Pattern becomes 1.000; "
        "if max CPS VoW is 0.903, that CPS also becomes 1.000."
    )
    st.markdown(
        "If a **Color by** choice expands the displayed points (e.g. CPS × Reservoir), "
        "the cross-plot uses the maximum of **those displayed points** and labels the "
        "effective aggregation. The 0–1 score is unitless, not a recovery factor; "
        "the raw VoW remains available in all ranking tables, hover and exports."
    )
    st.markdown("**3. Roll up to CPS properly.**")
    st.latex(r"\mathrm{VoW}_{\mathrm{CPS}}=\frac{\sum_{p\in\mathrm{CPS}} Q_{o,\mathrm{supported},p}}{\sum_{p\in\mathrm{CPS}} W_{\mathrm{inj},p}}")
    st.markdown("**Do not average the Pattern VoW ratios directly.** Summing volumes first avoids bias toward low-injection Patterns.")
    st.markdown("**4. Why can Value of Water exceed 2 (or even 4)?**")
    st.latex(r"\mathrm{VoW}=\frac{\mathrm{Oil\ Supported}}{\mathrm{Injected\ Water}}=\frac{432{,}000\ \mathrm{stb}}{100{,}000\ \mathrm{stb}}=4.32")
    st.markdown("VoW is a **ratio of two reported volumes**, not a percent, recovery factor, or bound between 0 and 1. "
                "Values above 2 are mathematically possible whenever reported Oil Supported is more than twice the injected-water volume over the same period. "
                "The oil may include pre-existing production, natural drive, pressure support from outside this Pattern, and a time-lagged response to earlier injection. "
                "Large values therefore warrant checking injection / production allocations, mismatched periods and zero / low injection — not automatic deletion.")
    high_vow = valid_pat.loc[valid_pat["VoW Technical"].gt(2)].sort_values("VoW Technical", ascending=False)
    st.caption(f"Current selection: {len(high_vow)} Patterns with Value of Water > 2.")
    if not high_vow.empty:
        st.dataframe(high_vow[["Pattern", "CPS", "Reservoir", "Manifold", "Supported", "Injection", "VoW Technical", "Supported mbd"]],
                     use_container_width=True, hide_index=True)
    st.markdown("**5. CPS and Manifold aggregation.**")
    st.latex(r"\mathrm{VoW}_{\mathrm{Manifold}}=\frac{\sum_{p \in \mathrm{Manifold}} Q_{o,\mathrm{supported},p}}{\sum_{p \in \mathrm{Manifold}} W_{\mathrm{inj},p}}")
    st.markdown("The uploaded **Manifold** column is read automatically. The Manifold name itself (e.g. CPS-2/FM09) is the physical grouping key: "
                "CPS parsed from Pattern names is NOT used to split a Manifold. Any Pattern CPS / Manifold CPS mismatch appears in QC. "
                "Unmapped rows remain in CPS/Pattern totals but are excluded from Manifold comparison. "
                "Patterns assigned to multiple physical Manifolds require supported allocation weights; duplicating their full volumes would overcount.")
    st.markdown("**6. Total, monthly average and daily rate (mbd).**")
    st.latex(r"Q_{\mathrm{avg/month}}=\frac{Q_{\mathrm{total\ period}}}{N_{\mathrm{months}}},\qquad\frac{Q_{o}/N}{W_{\mathrm{inj}}/N}=\frac{Q_o}{W_{\mathrm{inj}}}")
    st.latex(r"Q_{\mathrm{avg,day}}[\mathrm{mbd}]=\frac{Q_{\mathrm{total}}[\mathrm{stb}]}{N_{\mathrm{days}}\times1000}")
    st.markdown(f"Current daily-rate denominator: **{n_days:,} days**. Oil Supported on the VoW cross-plot defaults to **mbd** (thousand stb/day), independently of the global display. "
                "Use the actual reporting-period day count. All rates and unit changes leave the VoW ratio unchanged.")
    st.markdown("**7. Watercut and VRR.** Watercut is recomputed as produced water / (produced oil + produced water). Reported VRR is read from the workbook. Reservoir-corrected VRR requires FVF and full reservoir voidage terms, so we do **not** fabricate CPS reservoir VRR from stock-tank volumes.")
    if not valid_pat.empty:
        sample = valid_pat.iloc[0]
        st.markdown("**Worked example from the current filters:**")
        st.code(f"Pattern: {sample['Pattern']}\nOil Supported: {sample['Supported']:,.1f} stb\nWater Injection: {sample['Injection']:,.1f} stb\nRaw VoW = {sample['Supported']:,.1f} / {sample['Injection']:,.1f} = {sample['VoW Technical']:.4f} stb/stb\nNormalised VoW = {sample['VoW Normalized']:.4f} (0–1 score)", language="text")
    st.markdown("**Limitations:** contemporaneous production/injection correlation cannot establish causation; lagged response, injector-producer allocation factors, underlying decline, pressure support, reservoir connectivity, and water cycling must be evaluated before making injection-rate decisions.")
    st.markdown("Useful technical background: [SPE — Mature Flood Surveillance Using Streamlines](https://jpt.spe.org/mature-flood-surveillance-using-streamlines) and [SPE — Role of Surveillance Plots in Diagnosing Waterfloods](https://jpt.spe.org/twa/role-surveillance-plots-diagnosing-waterfloods).")

# -----------------------------------------------------------------------------
# Export all selected rows and rollups with calculation assumptions / QC.
# -----------------------------------------------------------------------------
st.divider()
st.subheader("Export results")
pattern_export = pat.sort_values("VoW Technical", ascending=False, na_position="last").reset_index(drop=True)
pattern_export.insert(0, "Rank (VoW)", [i + 1 if pd.notna(v) else np.nan for i, v in enumerate(pattern_export["VoW Technical"])])
cps_export = cps.sort_values("VoW Technical", ascending=False, na_position="last").reset_index(drop=True)
cps_export.insert(0, "Rank (VoW)", [i + 1 if pd.notna(v) else np.nan for i, v in enumerate(cps_export["VoW Technical"])])
manifold_export = manifold.sort_values("VoW Normalized", ascending=False, na_position="last").reset_index(drop=True) if not manifold.empty else pd.DataFrame(columns=["CPS", "Manifold", "Patterns", "VoW Normalized", "VoW Technical", "Supported mbd", "Injection mbd"])
if not manifold_export.empty:
    manifold_export.insert(0, "Rank (VoW)", [i + 1 if pd.notna(v) else np.nan for i, v in enumerate(manifold_export["VoW Technical"])])
ex1, ex2, ex3 = st.columns(3)
with ex1:
    st.download_button("Download Pattern Ranking CSV", pattern_export.to_csv(index=False).encode("utf-8-sig"),
                       file_name="Pattern_VoW_Ranking.csv", mime="text/csv", use_container_width=True)
with ex2:
    st.download_button("Download CPS Ranking CSV", cps_export.to_csv(index=False).encode("utf-8-sig"),
                       file_name="CPS_VoW_Ranking.csv", mime="text/csv", use_container_width=True)
with ex3:
    try:
        xlsx_bytes = file_download_excel(pattern_export, cps_export, manifold_export, work, qc_table, method, n_months, n_days, share, exclude_zero_supported, norm_references)
        st.download_button("Download Full Excel + QC + Guideline", xlsx_bytes,
                           file_name="Pattern_Waterflood_VoW_Analysis.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    except Exception as exc:
        st.error(f"Excel export failed: {exc}")
if not manifold_export.empty:
    st.download_button("Download Manifold Ranking CSV", manifold_export.to_csv(index=False).encode("utf-8-sig"),
                       file_name="Manifold_VoW_Ranking.csv", mime="text/csv", use_container_width=True)
st.caption("Source workbook and optional Manifold mapping are read in memory for the current Streamlit session. Do not commit confidential field datasets to a public GitHub repository.")
